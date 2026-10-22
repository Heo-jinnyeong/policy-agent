"""Create a form draft in memory without persisting applicant information."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from docx import Document

from .catalog import _extract_text, describe_form


def clean_values(form: dict, submitted: dict) -> dict[str, str]:
    allowed = {field["key"]: field for field in form["fields"]}
    if not allowed:
        raise ValueError("이 양식의 입력 항목을 확인할 수 없습니다. 양식 설정 파일을 추가해 주세요.")
    if set(submitted) - set(allowed):
        raise ValueError("양식에 없는 입력 항목이 있습니다.")
    values = {}
    for key, field in allowed.items():
        value = submitted.get(key, "")
        if not isinstance(value, str):
            raise ValueError(f"{field['label']} 값은 문자열이어야 합니다.")
        value = value.strip()
        if field["required"] and not value:
            raise ValueError(f"{field['label']} 항목을 입력해 주세요.")
        if len(value) > (2000 if field["multiline"] else 160):
            raise ValueError(f"{field['label']} 값이 너무 깁니다.")
        if any(ord(char) < 32 and char not in "\n\t" for char in value):
            raise ValueError(f"{field['label']} 값에 허용되지 않는 문자가 있습니다.")
        values[key] = value
    return values


def _fill_zip_template(path: Path, values: dict[str, str]) -> bytes | None:
    """Replace explicit {{field}} markers while retaining the package layout."""
    output = BytesIO()
    found = set()
    with ZipFile(path) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename.endswith(".xml"):
                for key, value in values.items():
                    marker = ("{{" + key + "}}").encode("utf-8")
                    if marker in data:
                        data = data.replace(marker, escape(value).encode("utf-8"))
                        found.add(key)
            target.writestr(member, data)
    return output.getvalue() if found == set(values) else None


def _companion_docx(form: dict, path: Path, values: dict[str, str]) -> bytes:
    document = Document()
    document.add_heading(form["title"] + " 작성 초안", 0)
    document.add_paragraph("원본 양식: " + form["filename"])
    if form["source_url"]:
        document.add_paragraph("원본 출처: " + form["source_url"])
    document.add_paragraph("아래 내용을 원본 신청서와 대조하여 확인하세요. 서명과 제출은 포함하지 않습니다.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "작성 내용"
    for field in form["fields"]:
        cells = table.add_row().cells
        cells[0].text = field["label"]
        cells[1].text = values.get(field["key"], "")
    if path.suffix.lower() == ".hwp":
        text = _extract_text(path).strip()
        if text:
            document.add_heading("원본 양식 내용 참고", level=1)
            document.add_paragraph(text[:4000])
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_draft(path: Path, submitted: dict, root: Path) -> tuple[bytes, str, bool]:
    form = describe_form(path, root)
    values = clean_values(form, submitted)
    if form["fills_original_layout"]:
        package = _fill_zip_template(path, values)
        if package is not None:
            return package, form["draft_format"], True
    return _companion_docx(form, path, values), "docx", False
