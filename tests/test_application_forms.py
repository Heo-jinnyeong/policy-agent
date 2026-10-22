import asyncio
import base64
from io import BytesIO
from zipfile import ZipFile

import httpx
from docx import Document

import server
from application_forms.catalog import describe_form, list_forms, search_forms
from application_forms.draft import build_draft


def _docx_with_markers() -> bytes:
    document = Document()
    document.add_heading("부전공 신청서", 0)
    document.add_paragraph("성명 {{student_name}} / 학번 {{student_id}}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_template_preserves_layout_and_does_not_change_original(tmp_path):
    folder = tmp_path / "academic"
    folder.mkdir()
    path = folder / "부전공 신청서.docx"
    original = _docx_with_markers()
    path.write_bytes(original)
    form = describe_form(path, tmp_path)
    assert form["fills_original_layout"]
    assert {item["key"] for item in form["fields"]} == {"student_name", "student_id"}
    assert search_forms("부전공 신청하고 싶어요", root=tmp_path)[0]["id"] == form["id"]
    content, extension, filled = build_draft(
        path, {"student_name": "홍길동", "student_id": "20260001"}, tmp_path)
    assert (extension, filled) == ("docx", True)
    assert "홍길동" in " ".join(p.text for p in Document(BytesIO(content)).paragraphs)
    assert path.read_bytes() == original


def test_hwpx_template_and_plain_docx_companion(tmp_path):
    folder = tmp_path / "academic"
    folder.mkdir()
    hwpx = folder / "부전공 신청.hwpx"
    with ZipFile(hwpx, "w") as archive:
        archive.writestr("Contents/section0.xml", "<root><t>성명 {{student_name}}</t></root>")
    form = describe_form(hwpx, tmp_path)
    assert form["draft_format"] == "hwpx"
    content, extension, filled = build_draft(hwpx, {"student_name": "김학생"}, tmp_path)
    assert (extension, filled) == ("hwpx", True)
    with ZipFile(BytesIO(content)) as archive:
        assert "김학생" in archive.read("Contents/section0.xml").decode()

    plain = folder / "일반 신청서.docx"
    document = Document()
    document.add_paragraph("성명")
    document.save(plain)
    assert not describe_form(plain, tmp_path)["fills_original_layout"]
    companion, extension, filled = build_draft(plain, {"student_name": "김학생"}, tmp_path)
    assert (extension, filled) == ("docx", False)
    assert "김학생" in " ".join(cell.text for row in Document(BytesIO(companion)).tables[0].rows
                                    for cell in row.cells)


def test_form_search_matches_korean_topic_with_particle(tmp_path):
    folder = tmp_path / "academic"
    folder.mkdir()
    (folder / "출석인정 신청서.docx").write_bytes(_docx_with_markers())
    matches = search_forms("출석인정은 어떤 경우에 받을 수 있나요?", root=tmp_path)
    assert len(matches) == 1
    assert matches[0]["title"] == "출석인정 신청서"


def test_forms_api_upload_and_draft(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "FORMS_DIR", tmp_path)
    payload = base64.b64encode(_docx_with_markers()).decode()

    async def check():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app),
                                     base_url="http://test") as client:
            uploaded = await client.post("/forms/upload", json={"filename": "부전공 신청서.docx",
                "domain": "academic", "content_base64": payload})
            assert uploaded.status_code == 200, uploaded.text
            form = uploaded.json()["form"]
            assert len((await client.get("/forms", params={"domain": "academic"})).json()["forms"]) == 1
            assert (await client.post(f"/forms/{form['id']}/draft",
                                      json={"values": {"student_name": "홍길동"}})).status_code == 422
            draft = await client.post(f"/forms/{form['id']}/draft",
                json={"values": {"student_name": "홍길동", "student_id": "20260001"}})
            assert draft.status_code == 200
            assert draft.headers["cache-control"] == "no-store"
            assert "홍길동" in " ".join(p.text for p in Document(BytesIO(draft.content)).paragraphs)
            assert (await client.get(f"/forms/{form['id']}/original")).content == base64.b64decode(payload)

    asyncio.run(check())
    assert len(list_forms(root=tmp_path)) == 1
